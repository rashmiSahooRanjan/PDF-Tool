import os
import sys
import shutil
from pathlib import Path
import pytest

# Ensure project root is in sys.path
BASE_DIR = Path(__file__).resolve().parent.parent
if str(BASE_DIR) not in sys.path:
    sys.path.insert(0, str(BASE_DIR))

import fitz
from PIL import Image
from docx import Document
from pypdf import PdfWriter, PdfReader

from app import create_app
from config import TestingConfig

@pytest.fixture(scope="session")
def app():
    app = create_app(TestingConfig)
    return app

@pytest.fixture(scope="session")
def client(app):
    return app.test_client()

@pytest.fixture(scope="session", autouse=True)
def setup_test_environment(tmp_path_factory):
    # Setup test workspace
    test_dir = str(tmp_path_factory.mktemp("docuflex_tests"))
    yield test_dir
    # Teardown
    shutil.rmtree(test_dir, ignore_errors=True)

@pytest.fixture
def sample_pdf(tmp_path):
    """Generate a clean synthetic 2-page PDF document."""
    pdf_path = str(tmp_path / "sample.pdf")
    doc = fitz.open()
    
    # Page 1
    p1 = doc.new_page(width=595, height=842) # A4
    p1.insert_text((50, 80), "DocuFlex Synthetic Test Document", fontsize=18, fontname="helv", color=(0.1, 0.2, 0.5))
    p1.insert_text((50, 120), "This is sample paragraph text on page one.", fontsize=12, fontname="helv")
    # Draw simple table grid
    p1.draw_rect(fitz.Rect(50, 160, 400, 260), color=(0.2, 0.4, 0.8), width=1)
    p1.draw_line((50, 200), (400, 200), color=(0.2, 0.4, 0.8), width=1)
    p1.draw_line((200, 160), (200, 260), color=(0.2, 0.4, 0.8), width=1)
    p1.insert_text((60, 185), "Header 1", fontsize=12, fontname="helv")
    p1.insert_text((210, 185), "Header 2", fontsize=12, fontname="helv")
    p1.insert_text((60, 235), "Row 1 Data", fontsize=11, fontname="helv")
    p1.insert_text((210, 235), "Row 2 Data", fontsize=11, fontname="helv")
    
    # Page 2
    p2 = doc.new_page(width=595, height=842)
    p2.insert_text((50, 80), "Second Page Content", fontsize=16, fontname="helv")
    p2.insert_text((50, 120), "DocuFlex multi-page verification test.", fontsize=12, fontname="helv")
    
    doc.save(pdf_path)
    doc.close()
    return pdf_path

@pytest.fixture
def sample_docx(tmp_path):
    """Generate a synthetic Word document with text and a table."""
    docx_path = str(tmp_path / "sample.docx")
    doc = Document()
    doc.add_heading("DocuFlex Sample Word Document", level=1)
    doc.add_paragraph("This document contains synthetic test data.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "Value 1"
    table.cell(1, 1).text = "Value 2"
    doc.save(docx_path)
    return docx_path

@pytest.fixture
def sample_image(tmp_path):
    """Generate a synthetic RGB PNG test image."""
    img_path = str(tmp_path / "sample.png")
    img = Image.new("RGB", (200, 200), color=(73, 109, 137))
    img.save(img_path)
    return img_path
